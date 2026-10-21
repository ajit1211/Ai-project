"""Helpers shared by Core's tests. Not imported by application code."""

import io
import zipfile
from collections.abc import Callable, Sequence
from pathlib import Path
from typing import Any

from docx import Document as new_docx
from pypdf import PdfWriter

from nexa_common.ids import new_id
from nexa_core.identity.models import User
from nexa_core.organizations.models import Organization, OrganizationMembership, Role, Team

PASSWORD = "correct-horse-battery-staple"

UserFactory = Callable[..., User]
OrgFactory = Callable[..., Organization]


def add_member(
    organization: Organization,
    user: User,
    role_key: str,
    status: str = OrganizationMembership.Status.ACTIVE,
) -> OrganizationMembership:
    role = Role.objects.get(organization=organization, key=role_key)
    return OrganizationMembership.objects.create(
        organization=organization, user=user, role=role, status=status
    )


# --- Document fixtures: small, synthetic, built in memory ------------------------------------


def _pdf_string(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return escaped.encode("latin-1")


def make_pdf(pages: list[str | None], *, title: str | None = None) -> bytes:
    """A valid PDF (correct xref) with one Helvetica line per page; None makes a blank page."""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # the pages object, filled once the page numbers are known
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    page_numbers: list[int] = []
    for text in pages:
        page_number = len(objects) + 1
        page_numbers.append(page_number)
        if text is None:
            objects.append(
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 3 0 R >> >> >>"
            )
            continue
        content = b"BT /F1 12 Tf 72 720 Td (" + _pdf_string(text) + b") Tj ET"
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (page_number + 1)
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    kids = b" ".join(b"%d 0 R" % n for n in page_numbers)
    objects[1] = b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(page_numbers)
    info_number = None
    if title is not None:
        objects.append(b"<< /Title (" + _pdf_string(title) + b") /Author (Test Author) >>")
        info_number = len(objects)

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n" % (len(objects) + 1))
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    trailer = b"trailer\n<< /Size %d /Root 1 0 R" % (len(objects) + 1)
    if info_number is not None:
        trailer += b" /Info %d 0 R" % info_number
    out.write(trailer + b" >>\nstartxref\n%d\n%%%%EOF\n" % xref)
    return out.getvalue()


def encrypt_pdf(data: bytes, *, user_password: str, owner_password: str) -> bytes:
    writer = PdfWriter(clone_from=io.BytesIO(data))
    writer.encrypt(user_password=user_password, owner_password=owner_password)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(*, with_table: bool = True) -> bytes:
    document = new_docx()
    document.core_properties.title = "Platform notes"
    document.core_properties.author = "Test Author"
    document.add_heading("Platform notes", level=1)
    document.add_paragraph("NEXA stores documents privately.")
    document.add_heading("Limits", level=2)
    document.add_paragraph("Fifty megabytes per file", style="List Bullet")
    document.add_paragraph("Five hundred pages per PDF", style="List Bullet")
    if with_table:
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Format"
        table.cell(0, 1).text = "Status"
        table.cell(1, 0).text = "PDF"
        table.cell(1, 1).text = "=SUM(A1)"
        table.cell(2, 0).text = "CSV"
        table.cell(2, 1).text = "ready | now"
        document.add_paragraph("After the table.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return out.getvalue()


def indexed_document(
    organization: Organization,
    owner: User | None,
    text: str,
    *,
    filename: str = "notes.txt",
    visibility: str = "private",
    team: Team | None = None,
    pages: Sequence[tuple[int, int, int]] = (),
    title: str = "",
    index: bool = True,
    **extra: Any,
) -> Any:
    """A processed (and by default indexed) document without a file, an upload or a worker:
    the row, its extracted text and its chunks, exactly as the pipeline would leave them.
    `pages` are (page, char_start, char_end) spans for a PDF's text."""
    from nexa_core.documents.indexing import run_indexing
    from nexa_core.documents.models import Document, DocumentIndex, DocumentText

    fmt = Path(filename).suffix.lstrip(".").lower() or "txt"
    document_id = new_id()
    document = Document.objects.create(
        id=document_id,
        organization=organization,
        uploaded_by=owner,
        visibility=visibility,
        team=team,
        filename=filename,
        title=title,
        format=fmt,
        media_type="text/plain",
        size_bytes=len(text.encode("utf-8")),
        sha256="0" * 64,
        storage_key=f"{organization.id}/{document_id}.{fmt}",
        status="completed",
        word_count=len(text.split()),
        char_count=len(text),
        page_count=len(pages) or None,
        **extra,
    )
    DocumentText.objects.create(
        document=document,
        content=text,
        pages=[{"page": p, "char_start": s, "char_end": e} for p, s, e in pages],
        char_count=len(text),
    )
    if index:
        DocumentIndex.objects.create(
            document=document, organization=organization, generation=1, status="pending"
        )
        outcome = run_indexing(document.id, 1)
        assert outcome == "indexed", outcome  # noqa: S101 - a test helper
    return document
